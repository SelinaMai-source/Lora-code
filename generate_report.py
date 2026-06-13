import json
from docx import Document

doc = Document()
doc.add_heading('实验v3 报告', 0)

with open('/root/autodl-tmp/Lora-code/results/runs/paper_instrdialog_ours_full_s123_v4_sota_41/final_metrics.json') as f:
    metrics = json.load(f)

doc.add_heading('SOTA 41 Metrics', level=1)
for key, value in metrics['final'].items():
    if key.startswith('eval.'):
        doc.add_paragraph(f"{key}: {value}")

doc.save('/root/autodl-tmp/Lora-code/实验v3.docx')
